"""Generate client-facing AI Interview Module overview (SOW + live Student Portal)."""

from pathlib import Path

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

doc = Document()

section = doc.sections[0]
section.top_margin = Inches(0.9)
section.bottom_margin = Inches(0.9)
section.left_margin = Inches(1.0)
section.right_margin = Inches(1.0)

style = doc.styles['Normal']
style.font.name = 'Calibri'
style.font.size = Pt(11)
style.font.color.rgb = RGBColor(0x10, 0x22, 0x28)
style._element.rPr.rFonts.set(qn('w:eastAsia'), 'Calibri')
style.paragraph_format.space_after = Pt(8)
style.paragraph_format.line_spacing = 1.15

for level, size, color in [
    (1, 18, RGBColor(0x0E, 0x5C, 0x6B)),
    (2, 14, RGBColor(0x0E, 0x5C, 0x6B)),
    (3, 12, RGBColor(0x0A, 0x3F, 0x49)),
]:
    hs = doc.styles[f'Heading {level}']
    hs.font.name = 'Calibri'
    hs.font.size = Pt(size)
    hs.font.bold = True
    hs.font.color.rgb = color
    hs.paragraph_format.space_before = Pt(16 if level == 1 else 12)
    hs.paragraph_format.space_after = Pt(6)

TEAL = RGBColor(0x0E, 0x5C, 0x6B)
INK = RGBColor(0x10, 0x22, 0x28)
MUTED = RGBColor(0x53, 0x60, 0x64)


def set_cell_shading(cell, hex_color):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    shd = OxmlElement('w:shd')
    shd.set(qn('w:fill'), hex_color)
    shd.set(qn('w:val'), 'clear')
    tc_pr.append(shd)


def add_para(text, *, bold=False, italic=False, size=11, color=INK, space_after=8):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(space_after)
    run = p.add_run(text)
    run.bold = bold
    run.italic = italic
    run.font.size = Pt(size)
    run.font.color.rgb = color
    run.font.name = 'Calibri'
    return p


def add_bullets(items):
    for item in items:
        p = doc.add_paragraph(item, style='List Bullet')
        p.paragraph_format.space_after = Pt(4)
        for run in p.runs:
            run.font.name = 'Calibri'
            run.font.size = Pt(11)
            run.font.color.rgb = INK


def add_table(headers, rows):
    table = doc.add_table(rows=1 + len(rows), cols=len(headers))
    table.style = 'Table Grid'
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    for i, h in enumerate(headers):
        cell = table.rows[0].cells[i]
        cell.text = h
        set_cell_shading(cell, '0E5C6B')
        for p in cell.paragraphs:
            p.paragraph_format.space_after = Pt(2)
            for run in p.runs:
                run.bold = True
                run.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
                run.font.size = Pt(10)
                run.font.name = 'Calibri'
    for r_i, row in enumerate(rows):
        for c_i, val in enumerate(row):
            cell = table.rows[r_i + 1].cells[c_i]
            cell.text = val
            if r_i % 2 == 1:
                set_cell_shading(cell, 'F1F5F6')
            for p in cell.paragraphs:
                p.paragraph_format.space_after = Pt(2)
                for run in p.runs:
                    run.font.size = Pt(10)
                    run.font.name = 'Calibri'
                    run.font.color.rgb = INK
    doc.add_paragraph()
    return table


# ---- Cover ----
add_para('Project K · Quirri B2B', bold=True, size=12, color=TEAL, space_after=4)
add_para('AI Interview Module — Client Overview', bold=True, size=22, color=TEAL, space_after=6)
add_para('Phase 1 B2B embedded interview (Student Portal)', size=12, color=MUTED, space_after=4)
add_para(
    'Aligned to ProjectK_B2B_SOW_v1.0 and the live Student Portal integration',
    size=11,
    color=MUTED,
    space_after=4,
)
add_para('5 October 2026 · Confidential — for client review', size=10, color=MUTED, space_after=16)

add_para(
    'This document explains how the AI Interview module works for institutional (B2B) clients '
    'under the Phase 1 Statement of Work. It describes the student experience as integrated in '
    'the Quirri Student Portal today, maps each capability to the SOW, and clearly separates '
    'items that sit in a separate Complete Interview / B2C SOW.',
    space_after=12,
)

# ---- 1 ----
doc.add_heading('1. Purpose of this document', level=1)
add_bullets([
    'Give the client a clear, accurate picture of the B2B embedded AI Interview under ProjectK_B2B_SOW_v1.0.',
    'Describe the end-to-end student journey as implemented in the Student Portal.',
    'Show which SOW items are live, which are contracted but still pending backend epics, and which are out of Phase 1 scope.',
])

# ---- 2 ----
doc.add_heading('2. Scope under the B2B agreement', level=1)
doc.add_heading('2.1 What Phase 1 contracts', level=2)
add_para('Per SOW §4, §5.2–§5.3, §6.4.5, §6.5.2 and §7:', space_after=6)
add_bullets([
    'AI Interview module embedded inside the B2B Student Portal (not a separate billed student product).',
    'Available to final-year B2B students without a separate student subscription.',
    'Students take AI mock interview sessions with voice interaction and an avatar interviewer.',
    'Interviews may be assigned by Faculty / HOD for specific cohorts, or self-initiated by the student.',
    'After the session: interview score report with per-question feedback and a personalised improvement plan.',
    'College Admin: interview cohort analytics (readiness, completion, weak areas, at-risk flags) and downloadable student / cohort reports.',
    'Interview session results shared by email after completion.',
])
add_para(
    'SOW note (§6.4.5): “For Complete Interview Module separate SOW is maintained.” '
    'Richer commercial interview packaging beyond this embedded B2B slice is not Phase 1 scope.',
    italic=True,
    color=MUTED,
    space_after=10,
)

doc.add_heading('2.2 Explicitly out of Phase 1 (SOW §14 and related)', level=2)
add_bullets([
    'B2C standalone Interview Portal and subscription / billing (separate B2C SOW).',
    'Mobile apps (iOS / Android).',
    'Offline or downloaded interview content.',
    'Live faculty-to-student video classes.',
    'On-premise deployment and third-party LMS integrations.',
])

# ---- 3 ----
doc.add_heading('3. How the interview works (student journey)', level=1)
add_para(
    'The integrated Student Portal flow follows five steps. This matches the live Interviews experience.',
    space_after=8,
)

steps = [
    (
        'Prepare',
        'The student signs in to the Student Portal, opens Interviews, and uploads a resume '
        '(PDF, DOC or DOCX, up to 10 MB). They can keep more than one resume and mark which one '
        'is current for the next session.',
    ),
    (
        'Set up',
        'The student chooses the target role (position), interview mode (Mock or Full), difficulty, '
        'and experience band. The system may check role fit against the resume and ask the student '
        'to confirm or switch if the field looks mismatched.',
    ),
    (
        'Interview',
        'The student joins a live room. An AI interviewer conducts a real-time voice conversation. '
        'In Full mode, a video avatar interviewer is used when the video service is available; '
        'otherwise the student can continue with voice only or cancel.',
    ),
    (
        'Review',
        'When the session ends, scoring runs automatically. The student is taken to a report page '
        'showing overall score, skill dimensions, question-by-question feedback, and an improvement plan.',
    ),
    (
        'Repeat',
        'Completed sessions appear under My reports with score history. The student can start another '
        'practice interview at any time (within institutional entitlement rules).',
    ),
]
for title, body in steps:
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(4)
    r = p.add_run(f'{title} — ')
    r.bold = True
    r.font.color.rgb = TEAL
    r.font.name = 'Calibri'
    r2 = p.add_run(body)
    r2.font.name = 'Calibri'
    r2.font.color.rgb = INK

# ---- 4 ----
doc.add_heading('4. Features available in the Student Portal today', level=1)
add_para(
    'The following capabilities are integrated in the Quirri B2B Student Portal interview screens '
    '(practice setup, live room, and reports). They support the SOW “mock interview + score report” deliverable.',
    space_after=8,
)

doc.add_heading('4.1 Resume handling', level=2)
add_bullets([
    'Upload resume (PDF / DOC / DOCX, up to 10 MB).',
    'List of uploaded resumes; mark one as current for the next interview.',
    'Interview questions are grounded on the current resume and chosen role.',
])

doc.add_heading('4.2 Role fit check', level=2)
add_bullets([
    'Before start, the platform can compare resume domain vs selected role.',
    'If mismatched, the student sees a clear message, a suggested role, and may switch or continue with acknowledgment.',
])

doc.add_heading('4.3 Interview set-up options', level=2)
add_bullets([
    'Mock interview — shorter, focused practice session.',
    'Full interview — longer session with richer post-session report; video avatar when enabled.',
    'Difficulty: beginner, intermediate, advanced, adaptive.',
    'Experience bands: fresher (0–1) through expert (7+ years).',
])
add_para(
    'SOW wording refers to “AI mock interview sessions.” Mock mode is the primary Phase 1 practice path. '
    'Full mode is available in the integrated product for a longer practice experience under the same '
    'embedded B2B module.',
    italic=True,
    color=MUTED,
    space_after=8,
)

doc.add_heading('4.4 Live interview experience', level=2)
add_bullets([
    'Real-time voice conversation with the AI interviewer (LiveKit-based room).',
    'On-screen transcript of both sides of the conversation.',
    'Status cues while the interviewer is listening, thinking, or speaking.',
    'Microphone mute / unmute during the session.',
    'Video avatar interviewer in Full mode when the institution’s video service is available; '
    'graceful voice-only fallback prompt if not.',
    'Natural session close and scoring of answers given so far if the student ends mid-session '
    'after the interviewer has joined.',
    'If the student leaves before the interviewer starts, the attempt can be cancelled so it is '
    'not scored as a completed interview.',
])

doc.add_heading('4.5 Scored performance report', level=2)
add_bullets([
    'Overall score with readiness banding.',
    'Skill dimension scores (for example communication, structure, technical depth, confidence — '
    'as returned by the scoring service).',
    'Question-by-question feedback: question, answer summary, written feedback, and turn score.',
    'Written performance summary and personalised action / practice plan.',
    'Reports generate shortly after the session and are opened from My reports or the post-session hand-off.',
])

doc.add_heading('4.6 Interview history and progress', level=2)
add_bullets([
    'My reports list of completed interviews with mode, date, and score.',
    'Summary KPIs: completed count, average score, best score.',
    'Filter by Mock / Full.',
    'Simple score trend over recent interviews.',
])

# ---- 5 ----
doc.add_heading('5. Faculty, HOD and College Admin (SOW vs status)', level=1)
add_para(
    'These items are in the Phase 1 B2B SOW. Student self-practice and scoring are live; assignment '
    'and college-level analytics depend on backend epics still being connected.',
    space_after=8,
)
add_table(
    ['SOW item', 'Portal', 'Status in this integration'],
    [
        [
            'Create / manage interview assignments for final-year cohorts',
            'Faculty / HOD · College Admin',
            'SOW in scope — UI placeholder pending interview assignment API (EPIC-18)',
        ],
        [
            'Final-year eligibility for B2B interview entitlement',
            'Institution / Student',
            'SOW rule — enforce with academic year / cohort data when assignment APIs are live',
        ],
        [
            'Interview cohort analytics (readiness, completion, weak areas, at-risk)',
            'College Admin',
            'SOW in scope — analytics wiring follows live reporting APIs',
        ],
        [
            'Downloadable student and cohort interview reports',
            'College Admin',
            'SOW in scope — report export follows live reporting APIs',
        ],
        [
            'Email of interview results after completion',
            'Notifications',
            'SOW in scope — delivery depends on notification / ESP configuration',
        ],
        [
            'Self-initiated mock interview + score report',
            'Student',
            'Live in Student Portal',
        ],
    ],
)

# ---- 6 ----
doc.add_heading('6. Access, security and privacy', level=1)
add_bullets([
    'Students must sign in to the institutional Student Portal before using Interviews.',
    'Each student sees only their own resumes, sessions, and reports.',
    'Role-based access separates Super Admin, College Admin / Faculty, and Student portals (SOW §8).',
    'Institutional data isolation applies — no cross-institution access.',
    'Media and uploaded content are stored on the platform infrastructure used for the deployment '
    '(SOW tech stack references object storage for session media and uploads).',
])

# ---- 7 ----
doc.add_heading('7. Platform building blocks (non-technical summary)', level=1)
add_bullets([
    'Real-time interview room: LiveKit (or equivalent as configured).',
    'AI avatar / video interviewer service when Full mode video is enabled.',
    'Speech recognition and scoring pipeline after the session.',
    'Student Portal UI: Interviews (setup), live room, report detail, My reports.',
])

# ---- 8 ----
doc.add_heading('8. Summary for the client', level=1)
add_table(
    ['Area', 'In Phase 1 B2B SOW?', 'In Student Portal integration?'],
    [
        ['Interview tab in Student Portal', 'Yes', 'Yes'],
        ['Voice AI mock interview', 'Yes', 'Yes'],
        ['Avatar / video interviewer (when enabled)', 'Yes (avatar)', 'Yes (Full mode + voice fallback)'],
        ['Self-initiated practice', 'Yes', 'Yes'],
        ['Score report + per-question feedback + improvement plan', 'Yes', 'Yes'],
        ['Resume upload and role-based questioning', 'Supports SOW interview', 'Yes'],
        ['Mock and Full practice modes', 'Mock named in SOW', 'Yes (both modes)'],
        ['Faculty / HOD cohort assignments', 'Yes', 'Pending API (EPIC-18)'],
        ['College Admin interview cohort analytics & downloads', 'Yes', 'Pending live reporting APIs'],
        ['Email of results post-completion', 'Yes', 'Pending notification delivery'],
        [
            'Session recording playback for the student',
            'Not a named SOW UI deliverable',
            'Not offered as a student playback feature',
        ],
        ['Spoken summary (English / Tanglish)', 'Not in Phase 1 SOW', 'Not Phase 1 contracted'],
        ['B2C standalone interview portal / billing', 'Excluded (§14)', 'Out of scope'],
        ['Complete Interview Module (full commercial package)', 'Separate SOW', 'Out of Phase 1 agreement'],
    ],
)

# ---- 9 ----
doc.add_heading('9. Closing note', level=1)
add_para(
    'Phase 1 delivers an embedded interview practice experience inside the institutional Student Portal: '
    'final-year B2B students can prepare, take a voice (and where available avatar) interview, and receive '
    'a scored report with actionable feedback. Faculty assignment workflows and college-wide interview '
    'analytics remain part of the same SOW and will light up as their APIs are connected. Anything beyond '
    'this embedded B2B slice — including the Complete Interview Module and the B2C portal — requires the '
    'separate SOW or a Change Request.',
    space_after=12,
)
add_para(
    'Prepared for client review · Project K (Quirri B2B) · Knotopian',
    bold=True,
    size=10,
    color=TEAL,
)
add_para('Reference agreement: ProjectK_B2B_SOW_v1.0', size=10, color=MUTED)

out_downloads = Path(r'c:\Users\Arun K\Downloads\ProjectK_B2B_AI_Interview_Module_Client_Overview.docx')
out_repo = Path(r'e:\b2b_superAdmin_frontend\docs\ProjectK_B2B_AI_Interview_Module_Client_Overview.docx')
out_repo.parent.mkdir(parents=True, exist_ok=True)
doc.save(str(out_downloads))
doc.save(str(out_repo))
print('WROTE', out_downloads)
print('WROTE', out_repo)
print('size_bytes', out_downloads.stat().st_size)
